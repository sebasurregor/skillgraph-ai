"""Genera las hojas de vida de prueba de esta carpeta.

    python docs/cv_ejemplo/generar_cv_ejemplo.py

Los PDF se escriben con un generador mínimo en Python puro (sin dependencias);
los Word, con python-docx. Los datos personales son ficticios.
"""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

AQUI = Path(__file__).resolve().parent
ANCHO, ALTO = 612, 792          # carta, en puntos

NEGRO = (0, 0, 0)
BLANCO = (1, 1, 1)
NAVY = (0, 0.14, 0.26)
VERDE = (0.05, 0.73, 0.31)
GRIS = (0.35, 0.42, 0.5)


class PDF:
    """Generador mínimo de PDF: texto Helvetica y rectángulos de color."""

    def __init__(self):
        self.paginas = [[]]

    def _op(self, texto):
        self.paginas[-1].append(texto)

    def rect(self, x, y, ancho, alto, color):
        self._op("%.3f %.3f %.3f rg %.1f %.1f %.1f %.1f re f"
                 % (*color, x, ALTO - y - alto, ancho, alto))

    def texto(self, x, y, cadena, tam=10, negrita=False, color=NEGRO):
        crudo = cadena.encode("cp1252")
        escapado = "".join(
            "\\%03o" % b if b > 126 or b in (40, 41, 92) else chr(b) for b in crudo)
        self._op("BT %.3f %.3f %.3f rg /%s %d Tf %.1f %.1f Td (%s) Tj ET"
                 % (*color, "F2" if negrita else "F1", tam, x, ALTO - y, escapado))

    def parrafo(self, x, y, cadena, ancho, tam=10, color=NEGRO, salto=None):
        """Escribe texto con ajuste de línea aproximado. Devuelve la y siguiente."""
        salto = salto or tam + 4
        maximo = max(10, int(ancho / (tam * 0.5)))
        linea = ""
        for palabra in cadena.split():
            if len(linea) + len(palabra) + 1 > maximo:
                self.texto(x, y, linea, tam, color=color)
                y += salto
                linea = palabra
            else:
                linea = (linea + " " + palabra).strip()
        if linea:
            self.texto(x, y, linea, tam, color=color)
            y += salto
        return y

    def guardar(self, ruta):
        objetos = [
            b"<< /Type /Catalog /Pages 2 0 R >>",
            None,   # páginas: se completa abajo
            b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica "
            b"/Encoding /WinAnsiEncoding >>",
            b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica-Bold "
            b"/Encoding /WinAnsiEncoding >>",
        ]
        hijos = []
        for operaciones in self.paginas:
            contenido = "\n".join(operaciones).encode("latin-1")
            objetos.append(b"<< /Length %d >>\nstream\n%s\nendstream"
                           % (len(contenido), contenido))
            objetos.append(
                b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 %d %d] /Contents %d 0 R "
                b"/Resources << /Font << /F1 3 0 R /F2 4 0 R >> >> >>"
                % (ANCHO, ALTO, len(objetos)))
            hijos.append(len(objetos))
        objetos[1] = b"<< /Type /Pages /Kids [%s] /Count %d >>" % (
            " ".join("%d 0 R" % h for h in hijos).encode(), len(hijos))

        salida = bytearray(b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\n")
        posiciones = []
        for numero, cuerpo in enumerate(objetos, start=1):
            posiciones.append(len(salida))
            salida += b"%d 0 obj\n%s\nendobj\n" % (numero, cuerpo)
        inicio_xref = len(salida)
        salida += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objetos) + 1)
        for posicion in posiciones:
            salida += b"%010d 00000 n \n" % posicion
        salida += (b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n"
                   % (len(objetos) + 1, inicio_xref))
        Path(ruta).write_bytes(bytes(salida))


# ---- 1. Formato plano (PDF): campos rotulados, una sola columna ---------------

def cv_plano():
    pdf = PDF()
    y = 70
    pdf.texto(60, y, "HOJA DE VIDA", 16, negrita=True)
    y += 34
    for linea in ("Nombre completo: Camilo Zapata Ospina",
                  "Correo electrónico: camilo.zapata@correo.com",
                  "Celular: 310 987 6543",
                  "Ciudad: Medellín, Antioquia"):
        pdf.texto(60, y, linea, 11)
        y += 17

    def seccion(titulo, lineas):
        nonlocal y
        y += 14
        pdf.texto(60, y, titulo, 12, negrita=True)
        y += 18
        for linea in lineas:
            y = pdf.parrafo(60, y, linea, 490, 10.5)

    seccion("PERFIL", [
        "Bachiller responsable y puntual, con disponibilidad inmediata para trabajos de "
        "temporada, turnos rotativos y fines de semana."])
    seccion("EXPERIENCIA LABORAL", [
        "Auxiliar de cargue - Distribuidora El Paisa (noviembre de 2024 a enero de 2025)",
        "Carga y descarga de mercancía y alistamiento de pedidos en temporada navideña.",
        "Vendedor de mostrador - Tienda Don Luis (2023)",
        "Atención al cliente y apoyo en la organización del local."])
    seccion("EDUCACIÓN", [
        "Bachiller académico - Institución Educativa Javiera Londoño, 2022"])
    seccion("HABILIDADES", [
        "Trabajo en equipo, puntualidad, atención al cliente, carga y descarga, "
        "alistamiento de pedidos."])
    seccion("REFERENCIAS", ["Luis Fernando Arango, administrador. Teléfono fijo 604 444 1122."])
    pdf.guardar(AQUI / "cv_plano.pdf")


# ---- 2. Formato Harvard (Word): nombre centrado y secciones en mayúscula ------

def cv_harvard():
    doc = Document()
    estilo = doc.styles["Normal"]
    estilo.font.name = "Times New Roman"
    estilo.font.size = Pt(11)

    nombre = doc.add_paragraph()
    nombre.alignment = WD_ALIGN_PARAGRAPH.CENTER
    trazo = nombre.add_run("Ana María Ríos Bedoya")
    trazo.bold = True
    trazo.font.size = Pt(16)
    contacto = doc.add_paragraph("Medellín, Colombia | 300 123 4567 | ana.rios@correo.com")
    contacto.alignment = WD_ALIGN_PARAGRAPH.CENTER

    def seccion(titulo):
        parrafo = doc.add_paragraph()
        parrafo.add_run(titulo).bold = True

    seccion("EDUCACIÓN")
    doc.add_paragraph("Universidad EAFIT, Medellín — Ingeniería de Sistemas, 2019")
    doc.add_paragraph("Curso de Scrum Fundamentals, 2021")

    seccion("EXPERIENCIA PROFESIONAL")
    doc.add_paragraph("TechNova S.A.S., Medellín — Desarrolladora backend (2021 – actualidad)")
    doc.add_paragraph("Diseño de servicios con Python y Django sobre bases de datos SQL.",
                      style="List Bullet")
    doc.add_paragraph("Despliegue de contenedores con Docker en servidores Linux.",
                      style="List Bullet")
    doc.add_paragraph("Andes Analytics, Bogotá — Analista de datos (2019 – 2021)")
    doc.add_paragraph("Reportes y consultas para el área comercial.", style="List Bullet")

    seccion("HABILIDADES TÉCNICAS")
    doc.add_paragraph("Python (avanzado), Django (intermedio), SQL (avanzado), Docker, Git, "
                      "Scrum, Linux")

    seccion("IDIOMAS")
    doc.add_paragraph("Español nativo, inglés B2")
    doc.save(AQUI / "cv_harvard.docx")


# ---- 3. Dos columnas (PDF) ----------------------------------------------------

def cv_dos_columnas():
    pdf = PDF()
    pdf.texto(60, 70, "VALENTINA RESTREPO MEJÍA", 20, negrita=True, color=NAVY)
    pdf.texto(60, 92, "Estilista profesional", 12, color=GRIS)
    pdf.rect(60, 104, 492, 1.5, VERDE)

    # Columna izquierda
    x, y = 60, 140
    pdf.texto(x, y, "CONTACTO", 11, negrita=True, color=NAVY)
    for linea in ("Móvil: 320 555 7788", "valentina.restrepo@correo.com", "Envigado, Antioquia"):
        y += 16
        pdf.texto(x, y, linea, 9.5)
    y += 34
    pdf.texto(x, y, "APTITUDES", 11, negrita=True, color=NAVY)
    for linea in ("Corte de cabello (experto)", "Colorimetría (avanzado)", "Manicure y pedicure",
                  "Maquillaje", "Atención al cliente"):
        y += 16
        pdf.texto(x, y, linea, 9.5)
    y += 34
    pdf.texto(x, y, "IDIOMAS", 11, negrita=True, color=NAVY)
    pdf.texto(x, y + 16, "Español", 9.5)

    # Columna derecha
    x, y = 260, 140
    pdf.texto(x, y, "SOBRE MÍ", 11, negrita=True, color=NAVY)
    y = pdf.parrafo(x, y + 16, "Estilista con seis años de trabajo en salones de belleza, "
                    "cumplida con los horarios y cercana con las clientas.", 290, 9.5)
    y += 18
    pdf.texto(x, y, "TRAYECTORIA LABORAL", 11, negrita=True, color=NAVY)
    y += 16
    for linea in ("Studio Bruna, Medellín. Estilista sénior, 2021 a la fecha.",
                  "Cortes, color y asesoría de imagen para clientela fija.",
                  "Peluquería Glamour, Envigado. Auxiliar de estilismo, 2018 a 2021.",
                  "Lavado, cepillado y apoyo en tintes."):
        y = pdf.parrafo(x, y, linea, 290, 9.5)
    y += 18
    pdf.texto(x, y, "FORMACIÓN ACADÉMICA", 11, negrita=True, color=NAVY)
    y += 16
    for linea in ("Técnica en peluquería y estética, Academia Marcel France, 2018.",
                  "Bachiller, Colegio La Presentación, 2016."):
        y = pdf.parrafo(x, y, linea, 290, 9.5)
    pdf.guardar(AQUI / "cv_dos_columnas.pdf")


# ---- 4. Con diseño (PDF): barra lateral de color y rótulos distintos ----------

def cv_diseno():
    pdf = PDF()
    pdf.rect(0, 0, 200, ALTO, NAVY)                    # barra lateral
    pdf.rect(200, 0, 412, 120, (0.91, 0.96, 0.99))     # banda del encabezado
    pdf.rect(60, 50, 80, 80, (0.4, 0.76, 0.93))        # espacio de la foto

    pdf.texto(225, 62, "Jorge Iván Castaño Ruiz", 22, negrita=True, color=NAVY)
    pdf.texto(225, 88, "Guarda de seguridad", 12, color=GRIS)

    x, y = 28, 175
    pdf.texto(x, y, "Datos personales", 12, negrita=True, color=VERDE)
    for rotulo, valor in (("Teléfono", "301 444 9080"), ("E-mail", "jorge.castano@correo.com"),
                          ("Ciudad", "Bello, Antioquia")):
        y += 22
        pdf.texto(x, y, rotulo, 8.5, negrita=True, color=(0.62, 0.76, 0.85))
        y += 13
        pdf.texto(x, y, valor, 9.5, color=BLANCO)
    y += 40
    pdf.texto(x, y, "Competencias", 12, negrita=True, color=VERDE)
    for linea in ("Curso de vigilancia", "Control de acceso", "Manejo de CCTV",
                  "Reporte de novedades", "Primeros auxilios", "Puntualidad"):
        y += 17
        pdf.texto(x, y, "• " + linea, 9.5, color=BLANCO)

    x, y = 225, 160
    pdf.texto(x, y, "Experiencia", 13, negrita=True, color=NAVY)
    pdf.rect(x, y + 6, 40, 2, VERDE)
    y += 28
    for cargo, fechas, detalle in (
            ("Guarda de seguridad · Vigilancia Antioquia", "2020 – 2025",
             "Control de ingreso de visitantes y vehículos en unidad residencial."),
            ("Operador de medios tecnológicos · Seguridad Atlas", "2017 – 2020",
             "Monitoreo de cámaras y registro de novedades en la minuta.")):
        pdf.texto(x, y, cargo, 10.5, negrita=True)
        pdf.texto(x, y + 14, fechas, 9, color=VERDE)
        y = pdf.parrafo(x, y + 28, detalle, 350, 9.5) + 12
    y += 10
    pdf.texto(x, y, "Estudios", 13, negrita=True, color=NAVY)
    pdf.rect(x, y + 6, 40, 2, VERDE)
    y += 28
    for linea in ("Curso de fundamentación en vigilancia, Academia Andina de Seguridad, 2017.",
                  "Bachiller académico, I.E. Fernando Vélez, 2012."):
        y = pdf.parrafo(x, y, linea, 350, 9.5)
    pdf.guardar(AQUI / "cv_diseno.pdf")


# ---- 5. Con diseño en tabla (Word) ---------------------------------------------

def cv_tabla():
    doc = Document()
    tabla = doc.add_table(rows=1, cols=2)
    izquierda, derecha = tabla.rows[0].cells

    titulo = izquierda.paragraphs[0]
    trazo = titulo.add_run("Luz Dary Mosquera Palacios")
    trazo.bold = True
    trazo.font.size = Pt(15)
    trazo.font.color.rgb = RGBColor(0x00, 0x24, 0x43)
    izquierda.add_paragraph("Auxiliar de cocina")
    izquierda.add_paragraph("Contacto").runs[0].bold = True
    izquierda.add_paragraph("Phone: 314 222 6677")
    izquierda.add_paragraph("Mail: luzdary.mosquera@correo.com")
    izquierda.add_paragraph("Qué sé hacer").runs[0].bold = True
    for habilidad in ("Manipulación de alimentos", "Cocina caliente", "Panadería y repostería",
                      "Trabajo en equipo"):
        izquierda.add_paragraph(habilidad, style="List Bullet")

    derecha.paragraphs[0].add_run("Historia laboral").bold = True
    derecha.add_paragraph("Restaurante La Provincia, Medellín. Auxiliar de cocina, 2022 a 2025.")
    derecha.add_paragraph("Preparación de almuerzos ejecutivos y apoyo en la línea caliente.")
    derecha.add_paragraph("Panadería El Trigal, Quibdó. Ayudante de panadería, 2019 a 2021.")
    derecha.add_paragraph("Estudios realizados").runs[0].bold = True
    derecha.add_paragraph("Técnico en cocina, SENA, 2019.")
    derecha.add_paragraph("Curso de manipulación de alimentos, 2024.")
    doc.save(AQUI / "cv_tabla.docx")


# ---- 6. Escaneada (PDF sin texto): solo una imagen simulada -------------------

def cv_escaneado():
    pdf = PDF()
    pdf.rect(40, 40, 532, 712, (0.93, 0.93, 0.9))
    for fila in range(18):
        pdf.rect(70, 90 + fila * 34, 300 + (fila * 53) % 170, 9, (0.55, 0.55, 0.55))
    pdf.guardar(AQUI / "cv_escaneado.pdf")


if __name__ == "__main__":
    for generar in (cv_plano, cv_harvard, cv_dos_columnas, cv_diseno, cv_tabla, cv_escaneado):
        generar()
        print("generado:", generar.__name__)
